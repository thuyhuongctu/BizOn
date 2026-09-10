#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Sinh hai tài liệu nộp kèm hồ sơ đăng ký quyền tác giả (chương trình máy tính):

  1. BANG_KE_DIA_<TÊN>_SHA256.xlsx  — bảng kê mọi tệp của bản đóng băng, mỗi tệp một mã băm SHA-256
  2. BAN_IN_MA_NGUON_<TÊN>_15-20trang.pdf — bản in mã nguồn tiêu biểu, đúng 20 trang

Chạy (mặc định cho BizOn, tại thư mục gốc kho mã):
  python3 scripts/release/build_copyright_filing_package.py [thư-mục-đầu-ra]

Dùng cho tác phẩm khác: chỉ định --repo-dir, --commit, --title, --repo-url,
--print-files và --name. Mọi nội dung lấy tại đúng commit đóng băng — chạy lại
luôn cho cùng kết quả.
"""
import argparse
import hashlib
import subprocess
import sys
from pathlib import Path

FREEZE_COMMIT = "f38c47e45bd87a13d543d5c8df8a28f8358c26da"
FREEZE_DATE = "03/8/2026, 11:01:46 (giờ Việt Nam)"
WORK_TITLE = "BizOn – Bật Nghiệp: Hệ thống phần mềm mô phỏng kinh doanh phục vụ đào tạo khởi nghiệp"
AUTHORS = "Đồng tác giả: Đỗ Thùy Hương · Phan Anh Tú"
REPO_URL = "https://github.com/thuyhuongctu/BizOn"
NAME = "BIZON"
REPO_DIR = "."
# Ba tệp mã nguồn tiêu biểu cho bản in 15–20 trang (900 dòng)
PRINT_FILES = ["js/engine.js", "js/quiz-bank.js", "js/backend.js"]
HEADER_LINE = "BizOn – Bật Nghiệp · Bản in mã nguồn nộp kèm hồ sơ đăng ký quyền tác giả"

FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")


def git(*args, binary=False):
    result = subprocess.run(["git", "-C", REPO_DIR, *args], capture_output=True, check=True)
    return result.stdout if binary else result.stdout.decode("utf-8")


def collect_manifest():
    rows = []
    listing = git("ls-tree", "-r", "-z", FREEZE_COMMIT)
    entries = [e for e in listing.split("\0") if e]
    batch = subprocess.Popen(["git", "-C", REPO_DIR, "cat-file", "--batch"],
                             stdin=subprocess.PIPE, stdout=subprocess.PIPE)
    for entry in entries:
        meta, path = entry.split("\t", 1)
        obj = meta.split()[2]
        batch.stdin.write((obj + "\n").encode())
        batch.stdin.flush()
        header = batch.stdout.readline().decode().split()
        size = int(header[2])
        data = batch.stdout.read(size)
        batch.stdout.read(1)
        rows.append((path, size, hashlib.sha256(data).hexdigest()))
    batch.stdin.close()
    batch.wait()
    rows.sort(key=lambda r: r[0])
    return rows


def build_xlsx(rows, out_path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    thin = Border(*[Side(style="thin", color="B0B0B0")] * 4)
    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(name="Times New Roman", bold=True, color="FFFFFF", size=11)
    base = dict(name="Times New Roman", size=11)

    ws = wb.active
    ws.title = "TONG_HOP"
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 92
    info = [
        ("BẢNG KÊ TỆP NỘP KÈM HỒ SƠ ĐĂNG KÝ QUYỀN TÁC GIẢ", ""),
        ("Tên tác phẩm", WORK_TITLE),
        ("Loại hình", "Chương trình máy tính"),
        ("Tác giả", AUTHORS.replace("Đồng tác giả: ", "")),
        ("Bản đóng băng (commit)", FREEZE_COMMIT),
        ("Thời điểm đóng băng", FREEZE_DATE),
        ("Kho mã nguồn", REPO_URL),
        ("Số tệp trong bảng kê", len(rows)),
        ("Tổng dung lượng (byte)", sum(r[1] for r in rows)),
        ("Thuật toán băm", "SHA-256"),
        ("Cách kiểm chứng", "Tại kho mã nguồn, chạy: git show " + FREEZE_COMMIT[:7]
         + ":<đường-dẫn-tệp> rồi tính SHA-256, đối chiếu với cột D của trang BANG_KE_TEP."),
        ("Ghi chú", "Bảng kê lập tự động bằng scripts/release/build_copyright_filing_package.py; "
         "nội dung mỗi đĩa nộp kèm là toàn bộ các tệp liệt kê tại trang BANG_KE_TEP."),
    ]
    for i, (label, value) in enumerate(info, start=1):
        a = ws.cell(row=i, column=1, value=label)
        b = ws.cell(row=i, column=2, value=value)
        a.font = Font(bold=True, **base) if i > 1 else Font(name="Times New Roman", bold=True, size=14)
        b.font = Font(**base)
        b.alignment = Alignment(wrap_text=True, vertical="top")
    ws.cell(row=9, column=2).number_format = "#,##0"

    ws2 = wb.create_sheet("BANG_KE_TEP")
    headers = ["STT", "Đường dẫn tệp", "Kích thước (byte)", "Mã băm SHA-256"]
    widths = [8, 72, 18, 70]
    for c, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws2.cell(row=1, column=c, value=h)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin
        ws2.column_dimensions[get_column_letter(c)].width = w
    mono = Font(name="Courier New", size=10)
    for i, (path, size, digest) in enumerate(rows, start=2):
        for c, v in enumerate((i - 1, path, size, digest), start=1):
            cell = ws2.cell(row=i, column=c, value=v)
            cell.border = thin
            cell.font = mono if c == 4 else Font(**base)
        ws2.cell(row=i, column=3).number_format = "#,##0"
    ws2.freeze_panes = "A2"
    wb.save(out_path)


def build_pdf(out_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas

    pdfmetrics.registerFont(TTFont("DejaVu", str(FONT_DIR / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVu-Bold", str(FONT_DIR / "DejaVuSans-Bold.ttf")))
    # Liberation Mono có đủ dấu tiếng Việt (DejaVu Sans Mono thiếu một số ký tự như ổ, ệ)
    lib = Path("/usr/share/fonts/truetype/liberation")
    pdfmetrics.registerFont(TTFont("DejaVuMono", str(lib / "LiberationMono-Regular.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVuMono-Bold", str(lib / "LiberationMono-Bold.ttf")))

    width, height = A4
    margin = 18 * mm
    c = canvas.Canvas(str(out_path), pagesize=A4)
    page_no = 0

    sources = []
    for f in PRINT_FILES:
        data = git("show", f"{FREEZE_COMMIT}:{f}", binary=True)
        text = data.decode("utf-8")
        digest = hashlib.sha256(data).hexdigest()
        lines = text.splitlines()
        sources.append((f, digest, lines))
    total_lines = sum(len(s[2]) for s in sources)

    WRAP_LIMIT = 100

    def wrap(line, limit=WRAP_LIMIT):
        chunks = []
        while len(line) > limit:
            chunks.append(line[:limit])
            line = line[limit:]
        chunks.append(line)
        return chunks

    # Lượt 1: dựng danh sách hàng in (mỗi hàng: kiểu, nội dung) để chia đúng 19 trang nội dung
    rows_to_draw = []
    for f, digest, lines in sources:
        rows_to_draw.append(("header", f"═══ TỆP: {f} · {len(lines)} dòng · SHA-256: {digest[:32]}…"))
        for n, line in enumerate(lines, start=1):
            for k, chunk in enumerate(wrap(line.replace("\t", "    "))):
                prefix = f"{n:>4} │ " if k == 0 else "     ↳ "
                rows_to_draw.append(("code", prefix + chunk))
        rows_to_draw.append(("gap", ""))
    while rows_to_draw and rows_to_draw[-1][0] == "gap":
        rows_to_draw.pop()

    CONTENT_PAGES = 19
    TOTAL_PAGES = CONTENT_PAGES + 1
    rows_per_page = -(-len(rows_to_draw) // CONTENT_PAGES)

    def header_footer():
        c.setFillGray(0.25)
        right = f"Commit {FREEZE_COMMIT[:7]} · {FREEZE_DATE}"
        c.setFont("DejaVu", 7)
        c.drawRightString(width - margin, height - 12 * mm, right)
        avail = width - 2 * margin - pdfmetrics.stringWidth(right, "DejaVu", 7) - 4 * mm
        left = HEADER_LINE
        c.setFont("DejaVu", 8)
        while left and pdfmetrics.stringWidth(left, "DejaVu", 8) > avail:
            left = left[:-2] + "…" if not left.endswith("…") else left[:-2] + "…"
        c.drawString(margin, height - 12 * mm, left)
        c.setLineWidth(0.4)
        c.line(margin, height - 13.5 * mm, width - margin, height - 13.5 * mm)
        c.line(margin, 16 * mm, width - margin, 16 * mm)
        c.drawString(margin, 11 * mm, "Ký nháy của tác giả: .............................")
        c.drawRightString(width - margin, 11 * mm, f"Trang {page_no}/{TOTAL_PAGES}")
        c.setFillGray(0)

    # ---- Trang bìa ----
    page_no += 1
    c.setFont("DejaVu-Bold", 15)
    y = height - 55 * mm
    c.drawCentredString(width / 2, y, "BẢN IN MÃ NGUỒN TIÊU BIỂU")
    c.setFont("DejaVu", 12)
    c.drawCentredString(width / 2, y - 9 * mm, "Nộp kèm hồ sơ đăng ký quyền tác giả — chương trình máy tính")
    c.setFont("DejaVu-Bold", 12)
    title_words = WORK_TITLE.split()
    title_lines, cur = [], ""
    for w in title_words:
        if len(cur) + len(w) + 1 > 46:
            title_lines.append(cur)
            cur = w
        else:
            cur = (cur + " " + w).strip()
    title_lines.append(cur)
    for i, part in enumerate(title_lines):
        c.drawCentredString(width / 2, y - (22 + i * 6.5) * mm, part)
    c.setFont("DejaVu", 10.5)
    rows_info = [
        (AUTHORS, ""),
        ("Kho mã nguồn:", REPO_URL),
        ("Bản đóng băng (commit):", FREEZE_COMMIT),
        ("Thời điểm đóng băng:", FREEZE_DATE),
        (f"Nội dung: {total_lines} dòng mã nguồn nguyên văn từ {len(sources):02d} tệp cốt lõi", ""),
    ]
    yy = y - 55 * mm
    for label, value in rows_info:
        c.drawString(margin + 8 * mm, yy, label + ("  " + value if value else ""))
        yy -= 6.5 * mm
    yy -= 3 * mm
    c.setFont("DejaVu-Bold", 10.5)
    c.drawString(margin + 8 * mm, yy, "Ba tệp được in và mã băm SHA-256 (tính tại đúng commit đóng băng):")
    yy -= 7 * mm
    c.setFont("DejaVuMono", 8)
    for f, digest, lines in sources:
        c.drawString(margin + 8 * mm, yy, f"{f}  ({len(lines)} dòng)")
        c.drawString(margin + 8 * mm, yy - 4 * mm, f"  SHA-256: {digest}")
        yy -= 10 * mm
    c.setFont("DejaVu", 9)
    yy -= 4 * mm
    for t in ["Mỗi trang in có số dòng liên tục theo tệp gốc, số trang và chỗ ký nháy.",
              "Kiểm chứng: git show " + FREEZE_COMMIT[:7] + ":<tệp> cho ra đúng nội dung in tại đây."]:
        c.drawString(margin + 8 * mm, yy, t)
        yy -= 5.5 * mm
    header_footer()
    c.showPage()

    # ---- Các trang mã nguồn: chia đều rows_to_draw thành đúng 19 trang ----
    top_y = height - 18 * mm
    bottom_y = 20 * mm
    line_h = (top_y - bottom_y) / rows_per_page
    code_size = min(8.2, line_h - 2.2)

    for start in range(0, len(rows_to_draw), rows_per_page):
        page_no += 1
        header_footer()
        y = top_y - line_h
        for kind, text in rows_to_draw[start:start + rows_per_page]:
            if kind == "header":
                c.setFont("DejaVuMono-Bold", code_size + 0.6)
            else:
                c.setFont("DejaVuMono", code_size)
            if text:
                c.drawString(margin, y, text)
            y -= line_h
        c.showPage()
    c.save()
    assert page_no == TOTAL_PAGES, f"kỳ vọng {TOTAL_PAGES} trang, in ra {page_no}"
    return page_no


def main():
    global FREEZE_COMMIT, FREEZE_DATE, WORK_TITLE, AUTHORS, REPO_URL, NAME, REPO_DIR, PRINT_FILES, HEADER_LINE
    ap = argparse.ArgumentParser()
    ap.add_argument("out_dir", nargs="?", default="artifacts/ip-evidence")
    ap.add_argument("--repo-dir", default=REPO_DIR)
    ap.add_argument("--commit", default=FREEZE_COMMIT)
    ap.add_argument("--freeze-date", default=None, help="ghi đè mô tả thời điểm đóng băng")
    ap.add_argument("--title", default=None)
    ap.add_argument("--repo-url", default=None)
    ap.add_argument("--name", default=None, help="tên viết hoa dùng trong tên tệp, ví dụ ENQUIZ")
    ap.add_argument("--print-files", nargs="*", default=None)
    ap.add_argument("--header-line", default=None)
    args = ap.parse_args()
    REPO_DIR = args.repo_dir
    FREEZE_COMMIT = git("rev-parse", args.commit).strip()
    if args.title:
        WORK_TITLE = args.title
    if args.repo_url:
        REPO_URL = args.repo_url
    if args.name:
        NAME = args.name
    if args.print_files:
        PRINT_FILES = args.print_files
    if args.freeze_date:
        FREEZE_DATE = args.freeze_date
    if args.header_line:
        HEADER_LINE = args.header_line

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    rows = collect_manifest()
    print(f"Bản đóng băng {FREEZE_COMMIT[:7]}: {len(rows)} tệp, {sum(r[1] for r in rows):,} byte")
    xlsx = out_dir / f"BANG_KE_DIA_{NAME}_SHA256.xlsx"
    build_xlsx(rows, xlsx)
    print(f"Đã ghi {xlsx}")
    pdf = out_dir / f"BAN_IN_MA_NGUON_{NAME}_15-20trang.pdf"
    pages = build_pdf(pdf)
    print(f"Đã ghi {pdf} ({pages} trang)")


if __name__ == "__main__":
    main()
